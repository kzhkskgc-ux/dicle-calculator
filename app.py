"""
DI/cle_OGTT Calculator  Version 1.4
神戸大学臨床糖尿病グループ

機能:
  - 手動入力による1症例計算
  - Excelファイルアップロードによる複数症例一括計算
  - 結果のExcelファイルダウンロード
  - Van Cauter法の個別パラメータ（性別・身長・体重・年齢・病型）によるISR0計算

DI/cle: Sugimoto H et al. J Clin Endocrinol Metab 2023;108:3080-3089.
OGIS120: Mari A et al. Diabetes Care 2001;24:539-548.
ISR0: Van Cauter E et al. Diabetes 1992;41:368-377.
"""

import streamlit as st
import numpy as np
import matplotlib
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from matplotlib.ticker import MultipleLocator
import pandas as pd
from io import BytesIO

# ── Japanese font ─────────────────────────────────────────────────────────────
def _setup_jp_font():
    candidates = ["Noto Sans CJK JP","NotoSansCJK-Regular",
                  "Hiragino Sans","Hiragino Kaku Gothic Pro",
                  "IPAexGothic","IPAPGothic","MS Gothic","Yu Gothic"]
    available = {f.name for f in fm.fontManager.ttflist}
    for name in candidates:
        if name in available:
            matplotlib.rcParams["font.family"] = name
            return
    for f in fm.fontManager.ttflist:
        if any(k in f.fname.lower() for k in ["notosanscjk","ipagothic","ipaexgothic"]):
            matplotlib.rcParams["font.family"] = f.name
            return
_setup_jp_font()
matplotlib.rcParams["axes.unicode_minus"] = False

st.set_page_config(page_title="DI/cle Calculator", page_icon="🧮", layout="centered")

# ═══════════════════════════════════════════════════════════════════════════════
# CONSTANTS
# ═══════════════════════════════════════════════════════════════════════════════
P1,P2,P3,P4 = 650,325,81_300,132
P5,P6 = 652e-6,173
V_OGIS,G_CLAMP = 10_000,90.0

VAN_CAUTER_TYPES = {
    "normal": {"label": "normal", "short_hl": 4.95, "fraction": 0.76},
    "obese":  {"label": "obese",  "short_hl": 4.55, "fraction": 0.78},
    "niddm":  {"label": "niddm",  "short_hl": 4.52, "fraction": 0.78},
}

OGIS_REFS = {
    "Lean\n(正常・痩せ型)":      {"mean":440,"sem":16,"color":"#2e7d32","n":15},
    "Obese\n(肥満・正常耐糖能)": {"mean":362,"sem":11,"color":"#1565c0","n":38},
    "IGT\n(耐糖能異常)":         {"mean":302,"sem":17,"color":"#e65100","n":13},
    "T2DM\n(2型糖尿病)":         {"mean":239,"sem":7, "color":"#b71c1c","n":38},
}

TEMPLATES = {
    "― テンプレートを選択 ―": None,
    "Lean（正常・痩せ型）": {
        "age":40,"sex":"男性","height":170.0,"weight":62.0,"dose":75.0,"vc_type":"normal",
        "G0":88.0,"G30":145.0,"G60":155.0,"G90":128.0,"G120":105.0,
        "I0":5.0,"I30":45.0,"I60":60.0,"I90":55.0,"I120":42.0,
        "C0":1.2,"C30":4.5,"C60":5.8,"C90":5.2,"C120":4.0},
    "Obese（肥満・正常耐糖能）": {
        "age":45,"sex":"男性","height":168.0,"weight":92.0,"dose":75.0,"vc_type":"obese",
        "G0":97.0,"G30":170.0,"G60":180.0,"G90":155.0,"G120":130.0,
        "I0":15.0,"I30":85.0,"I60":110.0,"I90":95.0,"I120":80.0,
        "C0":2.8,"C30":8.5,"C60":10.2,"C90":9.5,"C120":8.0},
    "IGT（耐糖能異常）": {
        "age":52,"sex":"男性","height":165.0,"weight":85.0,"dose":75.0,"vc_type":"obese",
        "G0":100.0,"G30":185.0,"G60":200.0,"G90":178.0,"G120":165.0,
        "I0":18.0,"I30":75.0,"I60":100.0,"I90":85.0,"I120":78.0,
        "C0":3.0,"C30":8.0,"C60":10.5,"C90":9.8,"C120":9.2},
    "T2DM（2型糖尿病）": {
        "age":58,"sex":"男性","height":163.0,"weight":78.0,"dose":75.0,"vc_type":"niddm",
        "G0":190.0,"G30":260.0,"G60":290.0,"G90":290.0,"G120":280.0,
        "I0":20.0,"I30":45.0,"I60":62.0,"I90":65.0,"I120":70.0,
        "C0":3.5,"C30":6.5,"C60":8.0,"C90":8.5,"C120":9.0},
}

# Excel入力テンプレートの列定義
EXCEL_COLS = [
    "ID","年齢","性別","身長(cm)","体重(kg)","負荷量(g)","VanCauter分類",
    "G0","G30","G60","G90","G120",
    "I0","I30","I60","I90","I120",
    "C0","C30","C60","C90","C120",
]

# ═══════════════════════════════════════════════════════════════════════════════
# FUNCTIONS
# ═══════════════════════════════════════════════════════════════════════════════
def calc_bsa(h,w): return 0.02350*(h**0.42246)*(w**0.51456)

def _normalize_sex(sex):
    s = str(sex).strip().lower()
    if s in ["男性", "男", "m", "male", "man", "1"]:
        return "m"
    if s in ["女性", "女", "f", "female", "woman", "0"]:
        return "f"
    raise ValueError("性別は 男性/女性 または m/f で入力してください。")

def _normalize_vc_type(vc_type):
    s = str(vc_type).strip().lower()
    mapping = {
        "normal": "normal", "lean": "normal", "ngt": "normal", "標準": "normal", "正常": "normal",
        "obese": "obese", "obesity": "obese", "肥満": "obese",
        "niddm": "niddm", "t2dm": "niddm", "type2": "niddm", "dm": "niddm", "2型糖尿病": "niddm", "糖尿病": "niddm",
    }
    if s in mapping:
        return mapping[s]
    return "normal"

def calc_van_cauter_params(age, sex, height_cm, weight_kg, vc_type="normal"):
    """
    Van Cauter et al. の個別パラメータ式に基づくC-peptide kinetics。

    入力:
      age: years
      sex: 男性/女性 or m/f
      height_cm, weight_kg
      vc_type: normal / obese / niddm

    出力:
      BSA, volume_mL, short_hl, long_hl, fraction, k01, k12, k21

    実装の考え方:
      C-peptide濃度の二指数関数表現において、短半減期・長半減期・短成分比率から
      two-compartment modelの速度定数を復元する。ここでk01がaccessible compartment
      からの不可逆的消失速度定数であり、 basal ISR = C0 * V1 * k01 とする。
    """
    age = float(age)
    height_cm = float(height_cm)
    weight_kg = float(weight_kg)
    sex_code = _normalize_sex(sex)
    vc_key = _normalize_vc_type(vc_type)
    bsa = calc_bsa(height_cm, weight_kg)

    # Van Cauter population formula for accessible compartment serum volume (mL)
    # BSA formula here follows the Van Cauter implementation commonly used with this model.
    bsa_vc = (weight_kg ** 0.425) * (height_cm ** 0.725) * 71.84 * 10**-4
    if sex_code == "m":
        volume_ml = (1.92 * bsa_vc + 0.64) * 1000
    else:
        volume_ml = (1.11 * bsa_vc + 2.04) * 1000

    short_hl = VAN_CAUTER_TYPES[vc_key]["short_hl"]
    long_hl = 0.14 * age + 29.2
    fraction = VAN_CAUTER_TYPES[vc_key]["fraction"]

    alpha = np.log(2) / short_hl
    beta = np.log(2) / long_hl
    A = fraction / volume_ml
    B = (1 - fraction) / volume_ml

    # Same algebra as the standard deconvolution implementation:
    k12 = (A * beta + alpha * B) / (A + B)
    k01 = alpha * beta / k12
    k21 = alpha + beta - (k12 + k01)

    return {
        "vc_type": vc_key,
        "bsa": bsa,
        "bsa_vc": bsa_vc,
        "volume_ml": volume_ml,
        "short_hl": short_hl,
        "long_hl": long_hl,
        "fraction": fraction,
        "k01": k01,
        "k12": k12,
        "k21": k21,
    }

def calc_ogis120(G0,G90,G120,I0,I90,dose,bsa):
    DO=dose/bsa; dG=(G120-G90)/30; dI=I90-I0+P2
    Cl=P4*(P1*DO-V_OGIS*dG)/G90/dI + P4*P3/G0/dI
    B=(P5*(G90-G_CLAMP)+1)*Cl
    disc=B**2+4*P5*P6*(G90-G_CLAMP)*Cl
    return Cl,(B+np.sqrt(max(disc,0)))/2,DO

def calc_isr0(C0_ng, age, bsa, sex="男性", height_cm=None, weight_kg=None, vc_type="normal"):
    """Basal ISR by individualized Van Cauter C-peptide kinetics.

    C0 is entered as ng/mL. 1 ng/mL C-peptide ≈ 0.3311 pmol/mL.
    Basal ISR (pmol/min) = C0(pmol/mL) * V1(mL) * k01(min^-1).
    The value is normalized by BSA to pmol/min/m².
    """
    if height_cm is None or weight_kg is None:
        # Backward-compatible fallback: use BSA-normalized apparent volume of 2.65 L/m².
        # New calculations should pass sex, height, and weight.
        k01 = 0.055
        return k01 * C0_ng * 331.1 * 2.65
    pars = calc_van_cauter_params(age, sex, height_cm, weight_kg, vc_type)
    c0_pmol_per_ml = C0_ng * 0.3311
    isr_total = c0_pmol_per_ml * pars["volume_ml"] * pars["k01"]
    isr_per_m2 = isr_total / pars["bsa"]
    return isr_per_m2

def calc_dicle(ogis,ISR0,G0,I0):
    """
    DI/cle_OGTT = ksen_OGTT * ksec / kcle^2

    Supplementary Appendix 8 defines:
      ksec = fasting ISR / fasting glucose level
      kcle = fasting ISR / fasting insulin level
      ksen_OGTT = OGIS

    Important:
    The source definition does not convert fasting glucose to mmol/L or insulin to pmol/L.
    OGIS formula itself uses glucose in mg/dL and insulin in microU/mL; therefore,
    the same clinical units are kept here for source-consistent DI/cle values.

    If G0 is converted to mmol/L and I0 to pmol/L, DI/cle becomes 648-fold larger
    (= 18 * 6^2), shifting log10(DI/cle) upward by log10(648) ≈ 2.812.
    """
    if G0<=0 or I0<=0 or ISR0<=0 or ogis<=0:
        return None,None,None
    ks=ISR0/G0      # ISR0 / fasting glucose level [mg/dL]
    kc=ISR0/I0      # ISR0 / fasting insulin level [microU/mL]
    return ogis*ks/(kc**2),ks,kc

def interp_ogis(v):
    if v>=400: return "正常域（Lean群相当）","#2e7d32","インスリン感受性は良好です（Lean群平均 440±16 ml·min⁻¹·m⁻²）。"
    elif v>=330: return "肥満正常群相当","#1565c0","インスリン感受性の軽度低下が示唆されます（Obese群平均 362±11）。"
    elif v>=270: return "IGT群相当","#e65100","インスリン抵抗性の進行が示唆されます（IGT群平均 302±17）。"
    else: return "2型糖尿病群相当","#b71c1c","著明なインスリン抵抗性が示唆されます（T2DM群平均 239±7）。"

def plot_ogis(val=None):
    fig,ax=plt.subplots(figsize=(7,3.2))
    fig.patch.set_facecolor("none"); ax.set_facecolor("#f8f8f5")
    for i,g in enumerate(OGIS_REFS):
        m,s,c=OGIS_REFS[g]["mean"],OGIS_REFS[g]["sem"],OGIS_REFS[g]["color"]
        ax.barh(i,m,color=c,alpha=0.2,height=0.55)
        ax.barh(i,m,color=c,alpha=0,height=0.55,
                xerr=s*2,error_kw=dict(ecolor=c,capsize=5,elinewidth=1.5,capthick=1.5))
        ax.text(m+s*2+8,i,f"{m} +/-{s}",va="center",fontsize=9,color=c,fontweight="bold")
    if val:
        ax.axvline(val,color="#c62828",lw=2.5,ls="--",zorder=5,label=f"今回: {val:.1f}")
        ax.legend(fontsize=9,framealpha=0.9,loc="lower right")
    ax.set_yticks(range(len(OGIS_REFS)))
    ax.set_yticklabels(list(OGIS_REFS.keys()),fontsize=9.5)
    ax.set_xlabel("OGIS120  (ml/min/m2)",fontsize=9)
    ax.set_title("OGIS120 参考値との比較  (Mari et al. 2001)",fontsize=9.5,pad=8)
    ax.set_xlim(0,580)
    ax.xaxis.set_minor_locator(MultipleLocator(50))
    ax.grid(axis="x",which="major",color="#ddd",lw=0.8)
    ax.grid(axis="x",which="minor",color="#eee",lw=0.4)
    ax.spines[["top","right"]].set_visible(False)
    ax.invert_yaxis(); fig.tight_layout(pad=1.0)
    return fig

def calc_one_row(row):
    """1症例分の計算。結果dictを返す。エラー時はNaN。"""
    try:
        age   = int(row["年齢"])
        sex   = row["性別"] if "性別" in row.index else "男性"
        h     = float(row["身長(cm)"])
        w     = float(row["体重(kg)"])
        dose  = float(row["負荷量(g)"])
        vc_type = row["VanCauter分類"] if "VanCauter分類" in row.index and pd.notna(row["VanCauter分類"]) else "normal"
        G0    = float(row["G0"]);  G90  = float(row["G90"]);  G120 = float(row["G120"])
        I0    = float(row["I0"]);  I90  = float(row["I90"])
        C0    = float(row["C0"])
        bsa   = calc_bsa(h,w)
        ClOGTT,ogis,DO = calc_ogis120(G0,G90,G120,I0,I90,dose,bsa)
        vcpars = calc_van_cauter_params(age, sex, h, w, vc_type)
        ISR0  = calc_isr0(C0, age, bsa, sex=sex, height_cm=h, weight_kg=w, vc_type=vc_type)
        dicle,k_sec,k_cle = calc_dicle(ogis,ISR0,G0,I0)
        log_d = float(np.log10(dicle)) if (dicle and dicle>0) else np.nan
        return {
            "BSA(m²)":     round(bsa,3),
            "BMI(kg/m²)":  round(w/(h/100)**2,1),
            "DO(g/m²)":    round(DO,1),
            "ClOGTT":      round(ClOGTT,1),
            "OGIS120":     round(ogis,1),
            "VanCauter分類_used": vcpars["vc_type"],
            "VC_V1(mL)": round(vcpars["volume_ml"],1),
            "VC_shortHL(min)": round(vcpars["short_hl"],2),
            "VC_longHL(min)": round(vcpars["long_hl"],2),
            "VC_fraction": round(vcpars["fraction"],2),
            "VC_k01(min^-1)": round(vcpars["k01"],5),
            "VC_k12(min^-1)": round(vcpars["k12"],5),
            "VC_k21(min^-1)": round(vcpars["k21"],5),
            "ISR0(pmol/min/m²)": round(ISR0,2),
            "k_sec":       round(k_sec,4),
            "k_cle":       round(k_cle,4),
            "DI/cle":      round(dicle,4) if dicle else np.nan,
            "log10(DI/cle)": round(log_d,3) if not np.isnan(log_d) else np.nan,
        }
    except Exception as e:
        return {k:np.nan for k in
                ["BSA(m²)","BMI(kg/m²)","DO(g/m²)","ClOGTT","OGIS120",
                 "VanCauter分類_used","VC_V1(mL)","VC_shortHL(min)","VC_longHL(min)","VC_fraction",
                 "VC_k01(min^-1)","VC_k12(min^-1)","VC_k21(min^-1)",
                 "ISR0(pmol/min/m²)","k_sec","k_cle","DI/cle","log10(DI/cle)"]}

def make_template_excel():
    """入力用テンプレートExcelを生成。"""
    # サンプル4行入り
    sample = [
        ["Lean-01",  40,"男性",170,62,75,"normal", 88,145,155,128,105, 5,45,60,55,42,  1.2,4.5,5.8,5.2,4.0],
        ["Obese-01", 45,"男性",168,92,75,"obese", 97,170,180,155,130,15,85,110,95,80,  2.8,8.5,10.2,9.5,8.0],
        ["IGT-01",   52,"男性",165,85,75,"obese",100,185,200,178,165,18,75,100,85,78,  3.0,8.0,10.5,9.8,9.2],
        ["T2DM-01",  58,"男性",163,78,75,"niddm",190,260,290,290,280,20,45,62,65,70,   3.5,6.5,8.0,8.5,9.0],
    ]
    df = pd.DataFrame(sample, columns=EXCEL_COLS)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="OGTT_Data")
        ws = writer.sheets["OGTT_Data"]
        # 列幅調整
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 12
    buf.seek(0)
    return buf

def make_result_excel(df_in, df_res):
    """入力＋結果を結合したExcelを生成。"""
    df_out = pd.concat([df_in.reset_index(drop=True),
                        df_res.reset_index(drop=True)], axis=1)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df_out.to_excel(writer, index=False, sheet_name="Results")
        ws = writer.sheets["Results"]
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 14
        # 結果列に背景色
        from openpyxl.styles import PatternFill
        fill = PatternFill(start_color="E8EEF8", end_color="E8EEF8", fill_type="solid")
        n_in = len(df_in.columns)
        for row in ws.iter_rows(min_row=2, min_col=n_in+1, max_col=len(df_out.columns)):
            for cell in row:
                cell.fill = fill
    buf.seek(0)
    return buf

# ═══════════════════════════════════════════════════════════════════════════════
# SESSION STATE
# ═══════════════════════════════════════════════════════════════════════════════
defs={"age":50,"sex":"男性","height":165.0,"weight":65.0,"dose":75.0,"vc_type":"normal",
      "G0":95.0,"G30":155.0,"G60":170.0,"G90":160.0,"G120":140.0,
      "I0":8.0,"I30":55.0,"I60":75.0,"I90":60.0,"I120":50.0,
      "C0":1.8,"C30":5.5,"C60":7.0,"C90":6.5,"C120":5.5}
for k,v in defs.items():
    if k not in st.session_state: st.session_state[k]=v

# ═══════════════════════════════════════════════════════════════════════════════
# UI
# ═══════════════════════════════════════════════════════════════════════════════

# ── Header ────────────────────────────────────────────────────────────────────
st.markdown("""
<div style="background:#1a3a6a;padding:18px 24px 14px;border-radius:10px;margin-bottom:6px">
  <div style="color:#90b4e8;font-size:0.78rem;font-weight:600;letter-spacing:.08em;margin-bottom:4px">神戸大学臨床糖尿病グループ</div>
  <div style="color:#fff;font-size:1.55rem;font-weight:700;line-height:1.2">DI/cle Calculator</div>
  <div style="color:#a8c4e8;font-size:0.82rem;margin-top:5px">Disposition Index / Clearance — OGTT-derived version</div>
  <div style="color:#f0c040;font-size:0.78rem;margin-top:8px;font-weight:500">Version 1.4　｜　ご使用は自己責任でお願いします</div>
</div>""",unsafe_allow_html=True)
st.caption(
    "DI/cle: Sugimoto H et al. *J Clin Endocrinol Metab* 2023; **108**: 3080–3089.  \n"
    "OGIS₁₂₀: Mari A et al. *Diabetes Care* 2001; **24**: 539–548.  \n"
    "ISR₀: Van Cauter E et al. *Diabetes* 1992; **41**: 368–377."
)
st.markdown("---")

# ── Reference ─────────────────────────────────────────────────────────────────
with st.expander("📊 OGIS₁₂₀ 参考値（クリックで表示）",expanded=False):
    st.dataframe(pd.DataFrame([
        {"群":g.replace("\n"," "),"mean":v["mean"],"±SEM":v["sem"],"n":v["n"],"単位":"ml·min⁻¹·m⁻²"}
        for g,v in OGIS_REFS.items()]),use_container_width=True,hide_index=True)
    st.caption("出典: Mari A et al. *Diabetes Care* 2001; **24**: 539–548. Table 3. ※ OGIS₁₈₀由来。")
    st.markdown("""<div style="display:grid;grid-template-columns:1fr 1fr;gap:8px;margin-top:8px">
      <div style="background:#e8f5e9;border-left:4px solid #2e7d32;padding:8px 12px;border-radius:4px"><b style="color:#2e7d32">≥400</b><span style="color:#2e7d32;font-size:0.88rem"> — 正常域</span></div>
      <div style="background:#e3f0fb;border-left:4px solid #1565c0;padding:8px 12px;border-radius:4px"><b style="color:#1565c0">330–399</b><span style="color:#1565c0;font-size:0.88rem"> — 肥満正常群</span></div>
      <div style="background:#fff3e0;border-left:4px solid #e65100;padding:8px 12px;border-radius:4px"><b style="color:#e65100">270–329</b><span style="color:#e65100;font-size:0.88rem"> — IGT群相当</span></div>
      <div style="background:#fce8e8;border-left:4px solid #b71c1c;padding:8px 12px;border-radius:4px"><b style="color:#b71c1c">＜270</b><span style="color:#b71c1c;font-size:0.88rem"> — T2DM群相当</span></div>
    </div>""",unsafe_allow_html=True)
st.markdown("---")

# ═══════════════════════════════════════════════════════════════════════════════
# MODE SELECTOR
# ═══════════════════════════════════════════════════════════════════════════════
mode = st.radio("入力モードを選択",
                ["📝 手動入力（1症例）", "📂 Excelアップロード（複数症例）"],
                horizontal=True)
st.markdown("---")

# ═══════════════════════════════════════════════════════════════════════════════
# MODE A: 手動入力
# ═══════════════════════════════════════════════════════════════════════════════
if mode == "📝 手動入力（1症例）":

    tk=st.selectbox("サンプルデータを読み込む（テンプレート）",list(TEMPLATES.keys()),index=0)
    if TEMPLATES[tk]:
        for k,v in TEMPLATES[tk].items(): st.session_state[k]=v
        st.info(f"「{tk}」のサンプルデータを読み込みました。値は自由に編集できます。")

    with st.expander("▶ 身体情報",expanded=True):
        c1,c2,c3,c4,c5=st.columns(5)
        age=c1.number_input("年齢",18,100,step=1,key="age")
        sex=c2.selectbox("性別",["男性","女性"],key="sex")
        height=c3.number_input("身長(cm)",100.0,220.0,step=0.1,key="height")
        weight=c4.number_input("体重(kg)",20.0,200.0,step=0.1,key="weight")
        dose=c5.number_input("負荷量(g)",50.0,100.0,step=1.0,key="dose")
        vc_type=st.selectbox(
            "Van Cauter分類（C-peptide kinetics）",
            ["normal", "obese", "niddm"],
            index=["normal", "obese", "niddm"].index(st.session_state.get("vc_type", "normal")),
            key="vc_type",
            help="short half-life と fraction を決める分類です。T2DMでは niddm、肥満例では obese を選択してください。"
        )
        bsa=calc_bsa(height,weight); bmi=weight/(height/100)**2
        vcpars_preview = calc_van_cauter_params(age, sex, height, weight, vc_type)
        m1,m2,m3,m4=st.columns(4)
        m1.metric("BSA (m²)",f"{bsa:.3f}",help="Gehan & George 1970")
        m2.metric("BMI (kg/m²)",f"{bmi:.1f}")
        m3.metric("VC V1 (mL)",f"{vcpars_preview['volume_ml']:.0f}")
        m4.metric("VC k01 (min⁻¹)",f"{vcpars_preview['k01']:.4f}")

    st.markdown("**OGTT 測定値**")
    _,th0,th30,th60,th90,th120=st.columns([2.2,1,1,1,1,1])
    for col,lab in zip([th0,th30,th60,th90,th120],["0 min","30 min","60 min","90 min","120 min"]):
        col.markdown(f"<div style='text-align:center;font-size:0.82rem;font-weight:600;color:#555'>{lab}</div>",unsafe_allow_html=True)

    gl,g0c,g30c,g60c,g90c,g120c=st.columns([2.2,1,1,1,1,1])
    gl.markdown("<div style='padding-top:8px;font-size:0.88rem;color:#444'>🩸 血糖 (mg/dl)</div>",unsafe_allow_html=True)
    G0  =g0c.number_input("G0",  label_visibility="collapsed",min_value=50.0,max_value=500.0,step=0.1,key="G0")
    G30 =g30c.number_input("G30", label_visibility="collapsed",min_value=50.0,max_value=600.0,step=0.1,key="G30")
    G60 =g60c.number_input("G60", label_visibility="collapsed",min_value=50.0,max_value=600.0,step=0.1,key="G60")
    G90 =g90c.number_input("G90", label_visibility="collapsed",min_value=50.0,max_value=600.0,step=0.1,key="G90")
    G120=g120c.number_input("G120",label_visibility="collapsed",min_value=50.0,max_value=600.0,step=0.1,key="G120")

    il,i0c,i30c,i60c,i90c,i120c=st.columns([2.2,1,1,1,1,1])
    il.markdown("<div style='padding-top:8px;font-size:0.88rem;color:#444'>💉 IRI (μU/ml)</div>",unsafe_allow_html=True)
    I0  =i0c.number_input("I0",  label_visibility="collapsed",min_value=0.0,max_value=300.0,step=0.1,key="I0")
    I30 =i30c.number_input("I30", label_visibility="collapsed",min_value=0.0,max_value=600.0,step=0.1,key="I30")
    I60 =i60c.number_input("I60", label_visibility="collapsed",min_value=0.0,max_value=600.0,step=0.1,key="I60")
    I90 =i90c.number_input("I90", label_visibility="collapsed",min_value=0.0,max_value=600.0,step=0.1,key="I90")
    I120=i120c.number_input("I120",label_visibility="collapsed",min_value=0.0,max_value=600.0,step=0.1,key="I120")

    cl2,c0c,c30c,c60c,c90c,c120c=st.columns([2.2,1,1,1,1,1])
    cl2.markdown("<div style='padding-top:8px;font-size:0.88rem;color:#444'>🔬 CPR (ng/ml)</div>",unsafe_allow_html=True)
    C0  =c0c.number_input("C0",  label_visibility="collapsed",min_value=0.0,max_value=20.0,step=0.01,key="C0")
    C30 =c30c.number_input("C30", label_visibility="collapsed",min_value=0.0,max_value=30.0,step=0.01,key="C30")
    C60 =c60c.number_input("C60", label_visibility="collapsed",min_value=0.0,max_value=30.0,step=0.01,key="C60")
    C90 =c90c.number_input("C90", label_visibility="collapsed",min_value=0.0,max_value=30.0,step=0.01,key="C90")
    C120=c120c.number_input("C120",label_visibility="collapsed",min_value=0.0,max_value=30.0,step=0.01,key="C120")
    st.caption("OGIS₁₂₀: G(0),G(90),G(120),I(0),I(90)を使用。ISR₀: C(0)とVan Cauter個別パラメータ（性別・身長・体重・年齢・分類）を使用。")

    if st.button("DI/cle を計算する",type="primary",use_container_width=True):
        errs=[]
        if I90-I0+P2<=0: errs.append("ΔI が非正です。")
        if C0<=0: errs.append("空腹時CPR C(0) が 0 以下です。")
        if I0<=0: errs.append("空腹時IRI I(0) が 0 以下です。")
        if errs:
            for e in errs: st.error(e)
            st.stop()

        bsa=calc_bsa(height,weight)
        ClOGTT,ogis120,DO=calc_ogis120(G0,G90,G120,I0,I90,dose,bsa)
        vcpars=calc_van_cauter_params(age, sex, height, weight, vc_type)
        ISR0=calc_isr0(C0, age, bsa, sex=sex, height_cm=height, weight_kg=weight, vc_type=vc_type)
        dicle,k_sec,k_cle=calc_dicle(ogis120,ISR0,G0,I0)
        log_dicle=np.log10(dicle) if dicle and dicle>0 else None
        olabel,ocolor,odesc=interp_ogis(ogis120)
        # Van Cauter parameters are computed individually from sex, age, height, weight, and subject type.

        st.markdown("---")
        st.markdown("### 結果")

        st.markdown("#### Step 1 — OGIS₁₂₀（k_sen）")
        r1,r2,r3=st.columns(3)
        r1.metric("ClOGTT（補正前）",f"{ClOGTT:.1f}",help="ml·min⁻¹·m⁻²")
        r2.metric("OGIS₁₂₀ = k_sen",f"{ogis120:.1f}",help="ml·min⁻¹·m⁻²")
        r3.metric("DO (g/m²)",f"{DO:.1f}")
        st.markdown(
            f"<div style='background:{ocolor}18;border-left:4px solid {ocolor};"
            f"padding:10px 14px;border-radius:6px;color:{ocolor};font-weight:600;margin:6px 0'>"
            f"OGIS₁₂₀ 判定：{olabel}<br>"
            f"<span style='font-weight:400;font-size:0.88rem'>{odesc}</span></div>",
            unsafe_allow_html=True)
        st.pyplot(plot_ogis(ogis120),use_container_width=True)
        st.caption("バー=mean、ひげ=±SEM×2、破線（赤）=今回の値。※ OGIS₁₈₀由来の参考値。")

        st.markdown("---")
        st.markdown("#### Step 2 — ISR₀（Van Cauter法）")
        s1,s2,s3=st.columns(3)
        s1.metric("C(0) (pmol/L)",f"{C0*331.1:.1f}")
        s2.metric("V1 (mL)",f"{vcpars['volume_ml']:.1f}",help="Van Cauter accessible compartment volume")
        s3.metric("ISR₀ (pmol/min/m²)",f"{ISR0:.2f}")
        p1,p2,p3,p4,p5=st.columns(5)
        p1.metric("分類",vcpars["vc_type"])
        p2.metric("short HL",f"{vcpars['short_hl']:.2f} min")
        p3.metric("long HL",f"{vcpars['long_hl']:.2f} min")
        p4.metric("fraction",f"{vcpars['fraction']:.2f}")
        p5.metric("k01",f"{vcpars['k01']:.5f}")
        st.caption(f"k12={vcpars['k12']:.5f} min⁻¹, k21={vcpars['k21']:.5f} min⁻¹。C0は ng/mL→pmol/mL に変換し、ISR0をBSAで補正しています。")

        st.markdown("---")
        st.markdown("#### Step 3 — DI/cle_OGTT")
        d1,d2,d3=st.columns(3)
        d1.metric("k_sec",f"{k_sec:.4f}",help="ISR₀/G₀ [G₀: mg/dL]。PDF Appendix 8の定義に合わせ、SI変換は行いません。")
        d2.metric("k_cle",f"{k_cle:.4f}",help="ISR₀/I₀ [I₀: μU/mL]。PDF Appendix 8の定義に合わせ、SI変換は行いません。")
        d3.metric("k_sen",f"{ogis120:.1f}")
        if dicle and dicle>0:
            e1,e2=st.columns(2)
            e1.metric("DI/cle_OGTT",f"{dicle:.4f}")
            e2.metric("log₁₀(DI/cle)",f"{log_dicle:.3f}")
            st.markdown(
                "<div style='background:#e8eef8;border-left:4px solid #1a4a8a;"
                "padding:12px 16px;border-radius:6px;color:#1a3a6a;margin-top:8px'>"
                "<b>DI/cle_OGTT の解釈</b><br>"
                "<span style='font-size:0.88rem;font-weight:400'>"
                "DI/cle は NGT と IGT の間で有意差なし、IGT と T2DM の間で有意差あり（Sugimoto et al. 2023）。<br>"
                "IGT では DI の低下をインスリンクリアランスの低下が代償し、DI/cle が保たれる可能性があります。"
                "</span></div>",unsafe_allow_html=True)

# ═══════════════════════════════════════════════════════════════════════════════
# MODE B: Excelアップロード
# ═══════════════════════════════════════════════════════════════════════════════
else:
    st.markdown("### Excelアップロード（複数症例一括計算）")

    # テンプレートダウンロード
    st.markdown("**① 入力用テンプレートをダウンロード**")
    st.download_button(
        label="📥 入力テンプレートをダウンロード (Excel)",
        data=make_template_excel(),
        file_name="DI_cle_input_template.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    st.caption(
        "テンプレートに従いデータを入力してください。  \n"
        "必須列: ID, 年齢, 性別, 身長(cm), 体重(kg), 負荷量(g), "
        "G0〜G120, I0〜I120, C0〜C120  \n"
        "性別: 男性 または 女性。VanCauter分類: normal / obese / niddm（省略時 normal）"
    )

    st.markdown("---")
    st.markdown("**② データファイルをアップロード**")
    uploaded = st.file_uploader("Excelファイルを選択 (.xlsx)", type=["xlsx"])

    if uploaded:
        try:
            df_in = pd.read_excel(uploaded)
            st.success(f"{len(df_in)} 症例を読み込みました。")
            st.dataframe(df_in, use_container_width=True)

            # 必須列チェック
            required = ["年齢","身長(cm)","体重(kg)","負荷量(g)",
                        "G0","G90","G120","I0","I90","C0"]
            missing = [c for c in required if c not in df_in.columns]
            if missing:
                st.error(f"以下の必須列がありません: {missing}")
                st.stop()

            # 一括計算
            results = [calc_one_row(row) for _,row in df_in.iterrows()]
            df_res = pd.DataFrame(results)

            st.markdown("---")
            st.markdown("### 計算結果")
            df_show = pd.concat([df_in[["ID"] if "ID" in df_in.columns else []].reset_index(drop=True),
                                  df_res], axis=1)
            st.dataframe(df_show, use_container_width=True)

            # Excelダウンロード
            st.markdown("**③ 結果をダウンロード**")
            st.download_button(
                label="📤 結果をダウンロード (Excel)",
                data=make_result_excel(df_in, df_res),
                file_name="DI_cle_results.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        except Exception as e:
            st.error(f"ファイルの読み込みに失敗しました: {e}")

# ── Formula ───────────────────────────────────────────────────────────────────
st.markdown("---")
with st.expander("計算式の詳細",expanded=False):
    st.markdown("""
#### Step 1: OGIS₁₂₀（Mari et al. 2001, Eq.8）
```
BSA     = 0.02350 × height^0.42246 × weight^0.51456
DO      = dose / BSA
ΔI      = I(90) − I(0) + 325
num     = 650·DO − 10000·[G(120)−G(90)]/30
ClOGTT  = 132·(num/G(90) + 81300/G(0)) / ΔI
B       = [6.52×10⁻⁴·(G(90)−90)+1]·ClOGTT
D       = B²+4·6.52×10⁻⁴·173·(G(90)−90)·ClOGTT
OGIS₁₂₀ = (B+√D)/2
```
#### Step 2: ISR₀（Van Cauter 1992：個別パラメータ）
```
BSA_VC = weight^0.425 × height^0.725 × 71.84 × 10^-4
V1     = (1.92×BSA_VC + 0.64)×1000  [男性]
       = (1.11×BSA_VC + 2.04)×1000  [女性]
short half-life = normal 4.95, obese 4.55, niddm 4.52 min
long half-life  = 0.14×age + 29.2 min
fraction        = normal 0.76, obese/niddm 0.78
α = ln(2)/short half-life
β = ln(2)/long half-life
A = fraction/V1
B = (1−fraction)/V1
k12 = (Aβ + αB)/(A+B)
k01 = αβ/k12
k21 = α + β − k12 − k01
C(0)[pmol/mL] = C(0)[ng/mL] × 0.3311
ISR₀[pmol/min/m²] = C(0)[pmol/mL] × V1[mL] × k01[min⁻¹] / BSA
```

#### Step 3: DI/cle_OGTT（Sugimoto et al. 2023）
```
k_sec = ISR₀ / G₀[mg/dL]
k_cle = ISR₀ / I₀[μU/mL]
DI/cle = OGIS₁₂₀ × k_sec / k_cle²

※ PDF Appendix 8の定義では、空腹時血糖値・空腹時インスリン値で割ると記載されており、
  G₀をmmol/L、I₀をpmol/Lへ変換するとは記載されていません。
  旧版のようにG₀/18、I₀×6を用いると、DI/cleは648倍、log₁₀(DI/cle)は約2.812高くなります。
```
""")

st.markdown("---")
st.caption(
    "神戸大学臨床糖尿病グループ　Version 1.4　｜　ご使用は自己責任でお願いします  \n"
    "Sugimoto H et al. *J Clin Endocrinol Metab* 2023; **108**: 3080–3089.  |  "
    "Mari A et al. *Diabetes Care* 2001; **24**: 539–548.  |  "
    "Van Cauter E et al. *Diabetes* 1992; **41**: 368–377."
)
